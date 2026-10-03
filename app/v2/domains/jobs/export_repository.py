"""Stream source observations into private, provenance-bearing exports."""
import csv
import io
import re
from datetime import date, datetime, time, timedelta, timezone
from typing import Any, Iterator

from openpyxl import Workbook
from sqlalchemy import select

from app.v2.common import REGULATORS
from app.v2.enforcement import scope
from app.v2.domains.auth.models import User
from app.v2.domains.collection.models import SourceFile
from app.v2.domains.enforcement.models import Enterprise, Inspection
from app.v2.domains.monitoring.models import Observation, Station
from app.v2.domains.monitoring.legacy import filtered


def safe_cell(value: Any) -> Any:
    # Prevent spreadsheet formula execution in text supplied by users/imports.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def export_job(db: Any, store: Any, parameters: dict[str, Any], owner_id: str | None, progress: Any) -> dict[str, Any]:
    import tempfile
    from pathlib import Path
    start, end = date.fromisoformat(parameters["start"]), date.fromisoformat(parameters["end"])
    regulatory = parameters.get("report") == "inspections"
    suffix = ".xlsx" if parameters.get("format") == "xlsx" else ".csv"
    count = 0
    source_hashes = set()
    with tempfile.TemporaryDirectory(dir=store.root / "exports") as directory, db.sessions() as session:
        output_path = Path(directory) / ("report" + suffix)
        user = session.get(User, owner_id)
        if not user or not user.active: raise ValueError("导出账号已失效")
        if regulatory:
            if user.role not in REGULATORS: raise ValueError("没有监管报告导出权限")
            columns = ["任务编号", "标题", "状态", "负责人编号", "期限", "核查事实", "版本"]
            created_start = datetime.combine(start, time.min, tzinfo=timezone.utc)
            created_end = datetime.combine(end + timedelta(days=1), time.min, tzinfo=timezone.utc)
            statement = scope(select(Inspection).where(Inspection.created_at >= created_start, Inspection.created_at < created_end), user)
            records: Iterator[list[Any]] = ([row.id,row.title,row.status,row.assignee_id,str(row.due_date or ""),row.facts,row.version]
                for row in session.scalars(statement.execution_options(yield_per=1000)))
        else:
            columns = ["站点", "观测时间UTC", "观测日期北京", "污染物", "浓度", "单位", "质量状态", "统计口径", "来源", "来源文件SHA256", "获取时间UTC"]
            observations_query: Any = filtered(start,end,parameters.get("station_id"),parameters.get("pollutant"),parameters.get("city_id"))
            statement = observations_query.with_only_columns(Station.name,Observation.observed_at,Observation.local_date,
                Observation.pollutant,Observation.value,Observation.unit,Observation.quality,Observation.basis,
                SourceFile.origin,SourceFile.sha256,SourceFile.created_at).order_by(Observation.observed_at,Station.name)
            def observation_records() -> Iterator[list[Any]]:
                for row in session.execute(statement.execution_options(yield_per=2000)):
                    if row[9]:
                        source_hashes.add(row[9])
                    yield [str(value) if value is not None else "" for value in row]
            records = observation_records()
        workbook = Workbook(write_only=True) if suffix == ".xlsx" else None
        with output_path.open("w",encoding="utf-8-sig",newline="") as output:
            if workbook is not None:
                append = workbook.create_sheet("数据").append
            else:
                append = csv.writer(output).writerow
            append(columns)
            for row in records:
                count += 1
                if count > 1000000: raise ValueError("单次导出超过一百万行，请缩小范围")
                append([safe_cell(value) for value in row])
                if count % 10000 == 0: progress(min(90, 10+count//10000))
        if workbook:
            meta = workbook.create_sheet("来源与说明")
            meta.append(["来源", "经授权的核查记录" if regulatory else "Quotsoft https://quotsoft.net/air/"])
            meta.append(["说明", "观测原始值；缺失不进行插补；不构成企业违法结论"])
            workbook.save(output_path)
        key, hashed = store.put_file(output_path,"exports",suffix)
    return {"storage_key":key,"sha256":hashed,"filename":f"{'核查汇总' if regulatory else '环境数据'}_{start}_{end}{suffix}",
        "rows":count,"source_hashes":sorted(source_hashes),"attribution":"业务核查记录" if regulatory else "https://quotsoft.net/air/","start":str(start),"end":str(end)}


def import_enterprises(db: Any, content: bytes, owner_id: str | None) -> dict[str, Any]:
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    required = {"企业名称", "统一社会信用代码", "行业", "地址", "联系人"}
    if not required.issubset(set(reader.fieldnames or [])): raise ValueError("企业CSV需包含：企业名称、统一社会信用代码、行业、地址、联系人")
    payload = list(reader)
    if len(payload) > 10000: raise ValueError("单次企业导入最多一万行")
    for row in payload:
        if not row["企业名称"].strip() or not re.fullmatch(r"[A-Z0-9]{18}", row["统一社会信用代码"].strip()):
            raise ValueError("企业名称为空或统一社会信用代码格式无效；本次未导入")
    inserted = skipped = 0
    with db.transaction() as session:
        from sqlalchemy import text
        session.execute(text("SELECT pg_advisory_xact_lock(hashtext('enterprise-import'))"))
        for row in payload:
            code = row["统一社会信用代码"].strip()
            if session.scalar(select(Enterprise.id).where(Enterprise.credit_code == code)):
                skipped += 1
                continue
            session.add(Enterprise(name=row["企业名称"].strip(), credit_code=code, industry=row["行业"], address=row["地址"], contact=row["联系人"], created_by=owner_id))
            session.flush()
            inserted += 1
    return {"inserted": inserted, "skipped": skipped}
