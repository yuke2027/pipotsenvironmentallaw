"""Property 26: exports and reports keep time, source, legal basis, and the disclaimer."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from hypothesis import given
from hypothesis import strategies as st
from openpyxl import load_workbook

from app.adapters.exports.common import ExportMeta, ExportRow
from app.core.fonts import ChineseFontStatus
from app.services.exports import export_dataset
from app.services.reports import build_report
from app.services.risk import DISCLAIMER
from tests.support.memory_auth import MemoryAuditRepository

_NOW = datetime(2026, 7, 1, tzinfo=timezone.utc)
_FONT_CANDIDATES = (
    Path(r"C:\Windows\Fonts\simhei.ttf"),
    Path(r"C:\Windows\Fonts\msyh.ttf"),
    Path(r"C:\Windows\Fonts\simsun.ttc"),
    Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJKSC-Regular.otf"),
)


def _font() -> ChineseFontStatus:
    for candidate in _FONT_CANDIDATES:
        if candidate.is_file():
            return ChineseFontStatus(True, candidate, candidate.stem, "ExportCJK", None, None)
    pytest.skip("no Chinese font is available for export rendering")
    raise AssertionError("unreachable")


def _text(path: Path, fmt: str) -> str:
    if fmt == "csv":
        return path.read_text(encoding="utf-8-sig")
    if fmt == "xlsx":
        workbook = load_workbook(path)
        pieces: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                pieces.append(" ".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(pieces)
    raw = path.read_bytes()
    parts = [raw.decode("latin-1", errors="ignore")]
    import re

    def _pdf_literal(body: bytes) -> str:
        output = bytearray()
        index = 0
        while index < len(body):
            if body[index] == 0x5C and index + 1 < len(body):
                code = body[index + 1]
                if index + 3 < len(body) and 48 <= code <= 55 and all(48 <= value <= 55 for value in body[index + 1 : index + 4]):
                    digits = body[index + 1 : index + 4]
                    output.append(int(digits.decode("ascii"), 8))
                    index += 1 + len(digits)
                    continue
                output.append(code)
                index += 2
                continue
            output.append(body[index])
            index += 1
        if output.startswith(b"\xfe\xff"):
            return output[2:].decode("utf-16-be", errors="ignore")
        return output.decode("latin-1", errors="ignore")

    for match in re.findall(rb"\((?:\\.|[^\\)])*\)", raw, flags=re.DOTALL):
        parts.append(_pdf_literal(match[1:-1]))
    return "\n".join(parts)


@given(
    fmt=st.sampled_from(("csv", "xlsx", "pdf")),
    source=st.sampled_from(("ONLINE", "CSV", "IMPUTED")),
    imputed=st.booleans(),
    penalty=st.one_of(st.none(), st.just("10-50万元")),
)
@pytest.mark.property
def test_property_26_exports_are_traceable(
    fmt: str,
    source: str,
    imputed: bool,
    penalty: str | None,
) -> None:
    """Files keep export time, query, source, legal version, and the penalty disclaimer.

    **Feature: environmental-supervision-platform, Property 26: 导出与报告可追溯**
    **Property 26：导出与报告可追溯**
    **Validates: Requirements 8.7, 11.7, 11.10, 12.1, 12.3, 12.4**
    """

    with TemporaryDirectory() as temporary:
        _assert_traceable(fmt, source, imputed, penalty, Path(temporary))


def _assert_traceable(fmt: str, source: str, imputed: bool, penalty: str | None, tmp_path: Path) -> None:
    font = _font() if fmt == "pdf" else None
    row = ExportRow(
        city="北京",
        station="东城",
        pollutant="PM2.5",
        value="80",
        source=source,
        is_imputed=imputed,
        legal_article="大气污染防治法第九十九条",
        standard_version="2012",
        penalty_estimate=penalty,
        measured_count=0 if imputed else 3,
        imputed_count=3 if imputed else 0,
        recommendation="加强监测",
    )
    meta = ExportMeta(
        exported_at=_NOW,
        timezone="Asia/Shanghai",
        query="city=北京",
        method="seasonal-median-robust-trend",
        period="2026-07",
    )
    audit = MemoryAuditRepository()
    destination = tmp_path / f"out.{fmt}"
    export_dataset(
        [row],
        fmt=fmt,
        meta=meta,
        destination=destination,
        font=font,
        actor_id=9,
        audit=audit,
        now=_NOW,
    )
    rendered = _text(destination, fmt)
    assert _NOW.isoformat() in rendered
    assert "city=北京" in rendered
    assert source in rendered
    assert ("插补" if imputed else "实测") in rendered
    assert "2012" in rendered
    assert "大气污染防治法第九十九条" in rendered
    if penalty:
        assert DISCLAIMER in rendered
    assert audit.entries[-1].action == "EXPORT_CREATED"
    assert audit.entries[-1].details is not None
    assert audit.entries[-1].details["format"] == fmt
    assert audit.entries[-1].details["user_id"] == "9"
    report_dest = tmp_path / f"report.{fmt if fmt != 'csv' else 'xlsx'}"
    report_fmt = "pdf" if fmt == "pdf" else "xlsx"
    build_report(
        [row],
        fmt=report_fmt,
        meta=meta,
        destination=report_dest,
        font=font,
        actor_id=9,
        audit=audit,
        now=_NOW,
    )
    report_text = _text(report_dest, report_fmt)
    assert "2026-07" in report_text
    assert DISCLAIMER in report_text
    assert "seasonal-median-robust-trend" in report_text
