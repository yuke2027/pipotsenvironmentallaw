"""Three-format export, report, and Chinese font fallback integration."""

from __future__ import annotations

import re
import shutil
import os
from datetime import datetime, timezone
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pytest
from openpyxl import load_workbook

from app.adapters.exports.common import ExportMeta, ExportRow
from app.core.errors import AppError, ErrorCode
from app.core.fonts import ChineseFontResolver
from app.services.exports import export_dataset
from app.services.reports import build_report
from app.services.risk import DISCLAIMER
from tests.support.memory_auth import MemoryAuditRepository

_NOW = datetime(2026, 8, 1, tzinfo=timezone.utc)
_FONT_CANDIDATES = (
    Path(r"C:\Windows\Fonts\simhei.ttf"),
    Path(r"C:\Windows\Fonts\msyh.ttf"),
    Path(r"C:\Windows\Fonts\simsun.ttc"),
    Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJKSC-Regular.otf"),
)


def _font_path() -> Path | None:
    configured = os.environ.get("CHINESE_FONT_PATH")
    candidates = ((Path(configured),) if configured else ()) + _FONT_CANDIDATES
    return next((path for path in candidates if path.is_file()), None)
_LAW = "大气污染防治法第九十九条"


def _pdf_text(path: Path) -> str:
    raw = path.read_bytes()
    parts = [raw.decode("latin-1", errors="ignore")]

    def literal(body: bytes) -> str:
        output = bytearray()
        index = 0
        while index < len(body):
            if body[index] == 0x5C and index + 3 < len(body) and all(48 <= value <= 55 for value in body[index + 1 : index + 4]):
                digits = body[index + 1 : index + 4]
                output.append(int(digits.decode("ascii"), 8))
                index += 1 + len(digits)
                continue
            if body[index] == 0x5C and index + 1 < len(body):
                output.append(body[index + 1])
                index += 2
                continue
            output.append(body[index])
            index += 1
        if output.startswith(b"\xfe\xff"):
            return output[2:].decode("utf-16-be", errors="ignore")
        return output.decode("latin-1", errors="ignore")

    for match in re.findall(rb"\((?:\\.|[^\\)])*\)", raw, flags=re.DOTALL):
        parts.append(literal(match[1:-1]))
    return "\n".join(parts)


def _resolver(root: Path, platform_name: str, family: str, candidates: tuple[str, ...]) -> ChineseFontResolver:
    source = _font_path()
    assert source is not None
    font_file = root / f"{platform_name}-cjk{source.suffix}"
    shutil.copy(source, font_file)

    def finder(_directories: object) -> list[Path]:
        return [font_file]

    return ChineseFontResolver(
        candidates,
        platform_name=platform_name,
        environment={"WINDIR": r"C:\Windows"},
        font_finder=finder,
        file_exists=lambda path: path == font_file,
        glyph_checker=lambda path: path == font_file,
        family_name_getter=lambda path: family,
    )


@pytest.mark.integration
def test_three_formats_reports_and_font_fallback(tmp_path: Path) -> None:
    if _font_path() is None:
        pytest.skip("a CJK font is required to render Chinese glyphs in this integration test")
    windows = _resolver(tmp_path, "Windows", "SimHei", ("Microsoft YaHei", "SimHei")).resolve()
    linux = _resolver(tmp_path, "Linux", "Noto Sans CJK SC", ("Microsoft YaHei", "Noto Sans CJK SC")).resolve()
    assert windows.available and linux.available
    assert windows.path is not None and linux.path is not None
    assert windows.path is not None and windows.path.name.startswith("Windows-cjk.")
    assert linux.path is not None and linux.path.name.startswith("Linux-cjk.")
    assert windows.matplotlib_font is not None
    assert Path(str(windows.matplotlib_font.get_file())) == windows.path

    row = ExportRow(
        city="北京",
        station="东城",
        pollutant="PM2.5",
        value="80",
        source="ONLINE",
        is_imputed=True,
        legal_article=_LAW,
        standard_version="2012",
        penalty_estimate="10-50万元",
        measured_count=2,
        imputed_count=1,
        recommendation="加强监测",
    )
    meta = ExportMeta(
        exported_at=_NOW,
        timezone="Asia/Shanghai",
        query="city=北京",
        method="seasonal-median-robust-trend",
        period="2026-08",
    )
    audit = MemoryAuditRepository()
    csv_path = export_dataset(
        [row],
        fmt="csv",
        meta=meta,
        destination=tmp_path / "out.csv",
        font=windows,
        actor_id=3,
        audit=audit,
        now=_NOW,
    )
    xlsx_path = export_dataset(
        [row],
        fmt="xlsx",
        meta=meta,
        destination=tmp_path / "out.xlsx",
        font=windows,
        actor_id=3,
        audit=audit,
        now=_NOW,
    )
    pdf_path = export_dataset(
        [row],
        fmt="pdf",
        meta=meta,
        destination=tmp_path / "out.pdf",
        font=linux,
        actor_id=3,
        audit=audit,
        now=_NOW,
    )
    csv_text = csv_path.read_bytes()
    assert csv_text.startswith(b"\xef\xbb\xbf")
    csv_body = csv_text.decode("utf-8-sig")
    assert "北京" in csv_body and "插补" in csv_body and _LAW in csv_body and "2012" in csv_body
    assert DISCLAIMER in csv_body and "Asia/Shanghai" in csv_body
    workbook = load_workbook(xlsx_path)
    sheet_text = " ".join(
        "" if cell is None else str(cell)
        for sheet in workbook.worksheets
        for excel_row in sheet.iter_rows(values_only=True)
        for cell in excel_row
    )
    assert "北京" in sheet_text and "插补" in sheet_text and _LAW in sheet_text and DISCLAIMER in sheet_text
    pdf_text = _pdf_text(pdf_path)
    assert "北京" in pdf_text and "插补" in pdf_text and _LAW in pdf_text and DISCLAIMER in pdf_text
    report = build_report(
        [row],
        fmt="pdf",
        meta=meta,
        destination=tmp_path / "report.pdf",
        font=windows,
        actor_id=3,
        audit=audit,
        now=_NOW,
    )
    report_text = _pdf_text(report)
    assert "2026-08" in report_text and "实测点数" in report_text and "插补点数" in report_text
    assert DISCLAIMER in report_text
    chart = tmp_path / "chart.png"
    figure, axis = plt.subplots()
    axis.set_title("实测与插补", fontproperties=windows.matplotlib_font)
    axis.plot([1, 2, 3], [2, 1, 4], label="历史")
    axis.legend(prop=windows.matplotlib_font)
    figure.savefig(chart)
    plt.close(figure)
    assert chart.stat().st_size > 100


@pytest.mark.integration
def test_unavailable_fonts_do_not_create_a_garbled_pdf(tmp_path: Path) -> None:
    resolver = ChineseFontResolver(
        ("SimHei", "Noto Sans CJK SC"),
        platform_name="Linux",
        font_finder=lambda _directories: [tmp_path / "empty.ttf"],
        file_exists=lambda _path: True,
        glyph_checker=lambda _path: False,
        family_name_getter=lambda _path: "SimHei",
    )
    status = resolver.resolve()
    assert status.available is False
    destination = tmp_path / "garbled.pdf"
    with pytest.raises(AppError) as caught:
        export_dataset(
            [
                ExportRow(
                    city="北京",
                    station="东城",
                    pollutant="PM2.5",
                    value="1",
                    source="CSV",
                    is_imputed=False,
                    legal_article=_LAW,
                    standard_version="2012",
                    penalty_estimate=None,
                )
            ],
            fmt="pdf",
            meta=ExportMeta(exported_at=_NOW, timezone="Asia/Shanghai", query="city=北京"),
            destination=destination,
            font=status,
            actor_id=1,
            audit=MemoryAuditRepository(),
            now=_NOW,
        )
    assert caught.value.code == ErrorCode.CHINESE_FONT_UNAVAILABLE
    assert not destination.exists()
